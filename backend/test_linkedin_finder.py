import os
import tempfile
import unittest

import pandas as pd
from openpyxl import load_workbook

from linkedin_finder import (Config, LinkedInFinder, Lookup, build_queries, process_file)


class FakeClient:
    """Returns canned organic_results; records queries."""
    def __init__(self, results):
        self.results, self.queries, self.calls = results, [], 0

    def search(self, q):
        self.queries.append(q)
        self.calls += 1
        return self.results if callable(self.results) is False else self.results(q)


def res(link, title, snippet=""):
    return {"link": link, "title": title, "snippet": snippet}


def lk(name, company="", city="", state="", title=""):
    return Lookup(row_index=0, full_name=name, company=company, city=city, state=state, title=title).prepare()


def run(results, lookup, **cfg):
    return LinkedInFinder(FakeClient(results), Config(**cfg)).find(lookup)


class Tests(unittest.TestCase):
    def test_high_confidence_match(self):
        r = run([res("https://www.linkedin.com/in/john-smith-123/",
                     "John Smith - Owner - Acme Distribution | LinkedIn",
                     "Location: New York · 500+ connections · Experience: Acme Distribution")],
                lk("John Smith", "Acme Distribution LLC", "New York", "NY", "Owner"))
        self.assertEqual(r.url, "https://www.linkedin.com/in/john-smith-123")
        self.assertEqual(r.confidence, "High")

    def test_wrong_company_stays_blank(self):
        r = run([res("https://www.linkedin.com/in/john-smith-9",
                     "John Smith - Manager - Globex | LinkedIn", "Location: New York")],
                lk("John Smith", "Acme Distribution", "New York"))
        self.assertEqual(r.url, "")
        self.assertTrue(r.status.startswith("Not verified"))

    def test_company_only_in_education_is_ignored(self):
        r = run([res("https://www.linkedin.com/in/john-smith-9", "John Smith - Student | LinkedIn",
                     "Education: Acme Distribution · Location: New York")],
                lk("John Smith", "Acme Distribution", "New York"))
        self.assertEqual(r.url, "")

    def test_ambiguous_is_blank(self):
        a = res("https://www.linkedin.com/in/john-smith-1", "John Smith - Owner - Acme | LinkedIn", "Location: New York")
        b = res("https://www.linkedin.com/in/john-smith-2", "John Smith - Owner - Acme | LinkedIn", "Location: New York")
        r = run([a, b], lk("John Smith", "Acme", "New York"))
        self.assertEqual(r.url, "")
        self.assertTrue(r.status.startswith("Ambiguous"))

    def test_location_conflict_goes_to_review(self):
        r = run([res("https://www.linkedin.com/in/john-smith-1", "John Smith - Owner - Acme | LinkedIn",
                     "Location: Austin, Texas")], lk("John Smith", "Acme", "New York", "NY"))
        self.assertEqual(r.url, "")
        self.assertEqual(len(r.review), 1)
        r2 = run([res("https://www.linkedin.com/in/john-smith-1", "John Smith - Owner - Acme | LinkedIn",
                      "Location: Austin, Texas")], lk("John Smith", "Acme", "New York", "NY"),
                 accept_location_conflict=True)
        self.assertNotEqual(r2.url, "")

    def test_nickname_is_medium(self):
        r = run([res("https://www.linkedin.com/in/mike-jones", "Mike Jones - CEO - Acme | LinkedIn", "")],
                lk("Michael Jones", "Acme"))
        self.assertEqual(r.confidence, "Medium")

    def test_different_person_same_company_rejected(self):
        r = run([res("https://www.linkedin.com/in/jane-smith", "Jane Smith - Owner - Acme | LinkedIn", "")],
                lk("John Smith", "Acme"))
        self.assertEqual(r.url, "")
        self.assertEqual(r.status, "Not found")

    def test_non_profile_links_ignored_and_url_canonical(self):
        r = run([res("https://www.linkedin.com/company/acme", "Acme | LinkedIn"),
                 res("https://in.linkedin.com/in/john-smith-5?trk=abc", "John Smith - Owner - Acme | LinkedIn", "")],
                lk("John Smith", "Acme"))
        self.assertEqual(r.url, "https://www.linkedin.com/in/john-smith-5")

    def test_no_results(self):
        r = run([], lk("John Smith", "Acme"))
        self.assertEqual(r.url, "")
        self.assertEqual(r.status, "Not found")

    def test_last_first_input_and_credentials(self):
        r = run([res("https://www.linkedin.com/in/john-smith", "John Smith, MBA - Owner - Acme | LinkedIn", "")],
                lk("Smith, John", "Acme"))
        self.assertTrue(r.url)

    def test_weak_company_needs_corroboration(self):
        c = [res("https://www.linkedin.com/in/john-smith", "John Smith - Owner | LinkedIn",
                 "Worked with Acme on projects")]
        self.assertEqual(run(c, lk("John Smith", "Acme Distribution")).url, "")

    def test_stops_early_to_save_credits(self):
        fc = FakeClient([res("https://www.linkedin.com/in/john-smith", "John Smith - Owner - Acme | LinkedIn", "")])
        LinkedInFinder(fc).find(lk("John Smith", "Acme", "New York"))
        self.assertEqual(len(fc.queries), 1)

    def test_queries(self):
        q = build_queries(lk("John Smith", "Acme Inc", "New York"), 3)
        self.assertEqual(q[0], 'site:linkedin.com/in "john smith" "Acme" "New York"')

    def test_missing_info_skipped(self):
        self.assertTrue(run([], lk("Cher", "Acme")).status.startswith("Skipped"))

    def test_end_to_end_excel(self):
        d = tempfile.mkdtemp()
        src = os.path.join(d, "in.csv")
        pd.DataFrame({"Decision Maker": ["John Smith", "Nobody Here", "John Smith"],
                      "Company Name": ["Acme", "Zzz Corp", "Acme"], "City": ["New York"] * 3,
                      "Phone": ["007", "1", "007"]}).to_csv(src, index=False)
        fc = FakeClient(lambda q: [res("https://www.linkedin.com/in/john-smith",
                                       "John Smith - Owner - Acme | LinkedIn", "Location: New York")]
                        if "john smith" in q else [])
        out = os.path.join(d, "out.xlsx")
        s = process_file(src, out, client=fc)
        ws = load_workbook(out)["Results"]
        hdr = [c.value for c in ws[1]]
        u = hdr.index("LinkedIn URL") + 1
        self.assertEqual(ws.cell(2, u).value, "https://www.linkedin.com/in/john-smith")
        self.assertIsNone(ws.cell(3, u).value)                       # blank, not a guess
        self.assertEqual(ws.cell(4, u).value, "https://www.linkedin.com/in/john-smith")
        self.assertEqual(ws.cell(2, hdr.index("Phone") + 1).value, "007")   # original data preserved
        self.assertEqual(s["found"], 2)


if __name__ == "__main__":
    unittest.main()

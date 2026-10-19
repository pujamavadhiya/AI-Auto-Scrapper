"""
LinkedIn decision-maker URL finder (SerpAPI) - converted from the n8n workflow
"Segment C - Decision Maker LinkedIn URL Finder (SerpAPI)".

Input : CSV / XLSX with decision-maker name, company, city (state/title optional)
Output: XLSX = your original columns + "LinkedIn URL" (BLANK when not verified)

Design rule: PRECISION OVER RECALL. A URL is written only when the Google result
itself shows evidence that it is the right person. Otherwise the cell stays blank
and the reason is written to "Match Status".

What is reused from the n8n workflow
  * SerpAPI Google search:  site:linkedin.com/in ...   (gl=us, hl=en, num=10)
  * Profile slug extraction + canonical URL  https://www.linkedin.com/in/<slug>
  * Title parsing  "Name - Title - Company | LinkedIn"  and the snippet parsing
    (Experience: ..., Location: ...)
  * Random delay between calls, de-duplication, "no result -> empty row"
What is new (because this is a lookup, not a discovery run)
  * One search plan per input row, built from name + company + city
  * Name / company / location verification, ambiguity detection, caching
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import random
import re
import sys
import time
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Optional

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Load environment variables from .env if present
try:
    from dotenv import load_dotenv
    load_dotenv()
    # Also check parent directory for .env
    load_dotenv(Path(__file__).resolve().parent.parent / ".env")
except ImportError:
    pass

# --------------------------------------------------------------------------- #
# Text helpers
# --------------------------------------------------------------------------- #

def strip_accents(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def norm_text(s) -> str:
    s = strip_accents(str(s or "")).lower().replace("&", " and ")
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s)).strip()


def has_phrase(haystack_norm: str, phrase_norm: str) -> bool:
    """Whole-word phrase search on already-normalised text."""
    if not phrase_norm:
        return False
    return re.search(rf"(?<![a-z0-9]){re.escape(phrase_norm)}(?![a-z0-9])", haystack_norm) is not None


# --------------------------------------------------------------------------- #
# Names
# --------------------------------------------------------------------------- #

NAME_NOISE = {"jr", "sr", "ii", "iii", "iv", "phd", "md", "mba", "cpa", "esq", "dds", "pmp",
              "cfa", "rn", "pe", "dr", "mr", "mrs", "ms", "prof", "mph", "jd", "cfp", "pmp"}

_NICK = [
    "michael mike mick mickey", "william will bill billy", "robert rob bob bobby robbie",
    "richard rick rich dick ricky", "james jim jimmy jamie", "john jack johnny jon",
    "joseph joe joey", "thomas tom tommy", "charles charlie chuck", "christopher chris",
    "daniel dan danny", "david dave", "matthew matt", "anthony tony", "steven stephen steve",
    "andrew andy drew", "edward ed eddie ted", "theodore ted theo", "nicholas nick",
    "benjamin ben", "jonathan jon", "timothy tim", "kenneth ken kenny", "gregory greg",
    "jeffrey jeff", "raymond ray", "ronald ron", "donald don", "lawrence larry",
    "zachary zach zack", "samuel sam", "alexander alex al", "jennifer jen jenny",
    "katherine catherine kathy kate katie kat", "margaret maggie meg peggy",
    "patricia pat patty trish", "deborah debbie deb", "susan sue susie", "rebecca becky",
    "samantha sam", "elizabeth liz beth betty", "victoria vicky", "abigail abby",
    "joshua josh", "peter pete", "philip phil", "patrick pat", "frederick fred freddy",
    "gerald gerry jerry", "douglas doug", "stanley stan", "vincent vince vinny",
    "salvatore sal", "nathaniel nathan nate", "mohammed muhammad mohammad mohamed",
]
NICK_GROUPS = [set(g.split()) for g in _NICK]


def first_names_match(a: str, b: str) -> Optional[str]:
    if a == b:
        return "exact"
    if any(a in g and b in g for g in NICK_GROUPS):
        return "nickname"
    return None


def name_tokens(raw: str, allow_swap: bool = False) -> list[str]:
    s = re.sub(r"\(.*?\)", " ", str(raw or ""))
    if "," in s:
        left, right = s.split(",", 1)
        right_t = norm_text(right).split()
        if allow_swap and right_t and len(right_t) <= 3 and not all(t in NAME_NOISE for t in right_t):
            s = f"{right} {left}"          # "Smith, John"  ->  "John Smith"
        else:
            s = left                        # "John Smith, MBA" -> "John Smith"
    return [t for t in norm_text(s).split() if t not in NAME_NOISE]


def name_match(inp: list[str], res: list[str]) -> str:
    """'exact' | 'nickname' | '' - first+last must agree, middle names ignored."""
    if len(inp) < 2 or len(res) < 2:
        return ""
    fi, li, fr, lr = inp[0], inp[-1], res[0], res[-1]
    last_ok = li == lr or (li in res[1:] and lr in inp[1:])      # compound surnames
    if not last_ok:
        return ""
    return first_names_match(fi, fr) or ""


# --------------------------------------------------------------------------- #
# Companies
# --------------------------------------------------------------------------- #

COMPANY_SUFFIX = {"inc", "incorporated", "llc", "ltd", "limited", "corp", "corporation", "co",
                  "company", "llp", "lp", "plc", "pllc", "pc", "gmbh", "pvt", "private", "the", "dba"}
COMPANY_GENERIC = {"group", "holdings", "holding", "international", "intl", "global", "services",
                   "service", "solutions", "enterprises", "enterprise", "industries", "systems",
                   "partners", "associates", "usa", "us", "america", "american", "national",
                   "worldwide", "trading", "distribution", "distributors", "distributor",
                   "wholesale", "brands", "brand", "management"}


def company_forms(raw: str) -> tuple[str, str]:
    """(full, core) normalised forms. full = without legal suffix; core = also without generic words."""
    toks = [t for t in norm_text(raw).split() if t not in COMPANY_SUFFIX]
    core = [t for t in toks if t not in COMPANY_GENERIC] or toks
    return " ".join(toks), " ".join(core)


def display_company(raw: str) -> str:
    return re.sub(r"[,\s]+(inc|llc|ltd|corp|co|corporation|company|l\.l\.c|incorporated)\.?$", "",
                  str(raw or "").strip(), flags=re.I).strip()


def company_strength(text_norm: str, full: str, core: str) -> str:
    """'strong' = full company phrase present; 'weak' = only distinctive core words; '' = none."""
    if not full:
        return ""
    if has_phrase(text_norm, full):
        return "strong"
    core_t = core.split()
    if core_t and all(has_phrase(text_norm, t) for t in core_t) and (len(core_t) >= 2 or len(core_t[0]) >= 4):
        return "weak"
    return ""


# --------------------------------------------------------------------------- #
# Locations
# --------------------------------------------------------------------------- #

US_STATES = {
    "AL": "alabama", "AK": "alaska", "AZ": "arizona", "AR": "arkansas", "CA": "california",
    "CO": "colorado", "CT": "connecticut", "DE": "delaware", "FL": "florida", "GA": "georgia",
    "HI": "hawaii", "ID": "idaho", "IL": "illinois", "IN": "indiana", "IA": "iowa", "KS": "kansas",
    "KY": "kentucky", "LA": "louisiana", "ME": "maine", "MD": "maryland", "MA": "massachusetts",
    "MI": "michigan", "MN": "minnesota", "MS": "mississippi", "MO": "missouri", "MT": "montana",
    "NE": "nebraska", "NV": "nevada", "NH": "new hampshire", "NJ": "new jersey", "NM": "new mexico",
    "NY": "new york", "NC": "north carolina", "ND": "north dakota", "OH": "ohio", "OK": "oklahoma",
    "OR": "oregon", "PA": "pennsylvania", "RI": "rhode island", "SC": "south carolina",
    "SD": "south dakota", "TN": "tennessee", "TX": "texas", "UT": "utah", "VT": "vermont",
    "VA": "virginia", "WA": "washington", "WV": "west virginia", "WI": "wisconsin", "WY": "wyoming",
    "DC": "district of columbia",
}
STATE_BY_NAME = {v: k for k, v in US_STATES.items()}


def resolve_state(raw: str) -> tuple[str, str]:
    """-> (abbr, full_name_norm); ('', '') if not a US state."""
    r = str(raw or "").strip()
    if r.upper() in US_STATES:
        return r.upper(), US_STATES[r.upper()]
    n = norm_text(r)
    if n in STATE_BY_NAME:
        return STATE_BY_NAME[n], n
    return "", n  # non-US region: keep normalised text for phrase matching


# --------------------------------------------------------------------------- #
# Query / candidate structures
# --------------------------------------------------------------------------- #

@dataclass
class Lookup:
    row_index: int
    full_name: str
    company: str = ""
    city: str = ""
    state: str = ""
    country: str = ""
    title: str = ""
    # derived
    tokens: list = field(default_factory=list)
    comp_full: str = ""
    comp_core: str = ""
    city_n: str = ""
    st_abbr: str = ""
    st_full: str = ""

    def prepare(self) -> "Lookup":
        self.tokens = name_tokens(self.full_name, allow_swap=True)
        self.comp_full, self.comp_core = company_forms(self.company)
        self.city_n = norm_text(self.city)
        self.st_abbr, self.st_full = resolve_state(self.state)
        return self

    def key(self) -> str:
        return "|".join([" ".join(self.tokens), self.comp_full, self.city_n, self.st_full, norm_text(self.title)])


@dataclass
class Candidate:
    url: str
    result_name: str
    headline: str
    snippet: str
    position: int
    name_kind: str = ""
    company: str = ""        # '', 'weak', 'strong'
    company_where: str = ""
    location: str = "unknown"  # match_city | match_state | conflict | unknown
    location_text: str = ""
    title_match: bool = False
    score: float = 0.0
    verdict: str = "reject"   # accept | review | reject
    confidence: str = ""
    reasons: list = field(default_factory=list)

    def evidence(self) -> str:
        bits = [f"name:{self.name_kind or 'no'}"]
        if self.company:
            bits.append(f"company:{self.company}({self.company_where})")
        if self.location != "unknown":
            bits.append(f"location:{self.location}" + (f"({self.location_text})" if self.location_text else ""))
        if self.title_match:
            bits.append("title:match")
        return "; ".join(bits)


@dataclass
class RowResult:
    url: str = ""
    status: str = "Not found"
    confidence: str = ""
    matched_name: str = ""
    headline: str = ""
    evidence: str = ""
    searches: int = 0
    review: list = field(default_factory=list)


# --------------------------------------------------------------------------- #
# Parsing a SerpAPI organic result (mirrors the n8n "Normalize and Filter" node)
# --------------------------------------------------------------------------- #

_LINK_RE = re.compile(r"^https?://(?:[a-z0-9-]+\.)?linkedin\.com/in/([^/?#\s]+)", re.I)
_TITLE_SYNONYMS = [("chief executive officer", "ceo"), ("chief operating officer", "coo"),
                   ("chief marketing officer", "cmo"), ("chief technology officer", "cto"),
                   ("chief financial officer", "cfo"), ("vice president", "vp"),
                   ("president", "president")]
_TITLE_STOP = {"of", "the", "and", "at", "for", "to", "in", "a", "an"}


def _title_norm(s: str) -> str:
    n = norm_text(s)
    for long, short in _TITLE_SYNONYMS:
        n = n.replace(long, short)
    return n


def parse_result(r: dict, position: int) -> Optional[Candidate]:
    m = _LINK_RE.match(str(r.get("link") or ""))
    if not m:
        return None                                   # not a /in/ profile
    url = f"https://www.linkedin.com/in/{m.group(1).rstrip('/')}"

    raw_title = str(r.get("title") or "")
    clean = re.sub(r"\s*[|\-\u2013\u2014]\s*LinkedIn\s*$", "", raw_title, flags=re.I)
    parts = [p.strip() for p in re.split(r"\s+[-\u2013\u2014]\s+|\s+\|\s+", clean) if p.strip()]
    if not parts:
        return None
    return Candidate(url=url, result_name=parts[0], headline=" - ".join(parts[1:]),
                     snippet=str(r.get("snippet") or ""), position=position)


def split_snippet(snippet: str) -> dict:
    """Split a Google LinkedIn snippet into experience / location / other text."""
    out = {"experience": "", "location": "", "rest": ""}
    rest = []
    for seg in re.split(r"[\u00b7\u2022|]", snippet):
        s = seg.strip()
        low = s.lower()
        if not s:
            continue
        if low.startswith("experience:"):
            out["experience"] = s.split(":", 1)[1].strip()
        elif low.startswith("education:"):
            continue                                   # never use education as company evidence
        elif low.startswith("location:"):
            out["location"] = s.split(":", 1)[1].strip()
        elif re.search(r"\b(connections?|followers?)\b", low):
            continue
        else:
            rest.append(s)
    if not out["location"]:                            # fallback: "New York, NY" / "Greater X Area"
        for s in rest[:3]:
            if re.fullmatch(r"[A-Z][A-Za-z.' -]+,\s*[A-Za-z .'-]+(,\s*United States)?", s) or s.endswith(" Area"):
                out["location"] = s
                break
    out["rest"] = " ".join(rest)
    return out


def eval_location(loc_text: str, lk: Lookup) -> str:
    if not loc_text or not (lk.city_n or lk.st_full):
        return "unknown"
    loc_n = norm_text(loc_text)
    if lk.city_n and has_phrase(loc_n, lk.city_n):
        return "match_city"
    if lk.st_full and has_phrase(loc_n, lk.st_full):
        return "match_state"
    if lk.st_abbr and re.search(rf"\b{lk.st_abbr}\b", loc_text):   # case-sensitive on purpose
        return "match_state"
    return "conflict"


def evaluate(c: Candidate, lk: Lookup, cfg: "Config") -> None:
    c.name_kind = name_match(lk.tokens, name_tokens(c.result_name))
    if not c.name_kind:
        c.verdict = "reject"
        return

    snip = split_snippet(c.snippet)
    head_n = norm_text(c.headline)
    exp_n = norm_text(snip["experience"])
    rest_n = norm_text(snip["rest"])

    # company evidence (strong only from headline or the "Experience:" field)
    if lk.comp_full:
        for where, text in (("headline", head_n), ("experience", exp_n)):
            s = company_strength(text, lk.comp_full, lk.comp_core)
            if s == "strong" or (s == "weak" and c.company != "strong"):
                if s == "strong" or not c.company:
                    c.company, c.company_where = s, where
            if c.company == "strong":
                break
        if not c.company:
            s = company_strength(rest_n, lk.comp_full, lk.comp_core)
            if s:
                c.company, c.company_where = "weak", "snippet"      # snippet body = weak at best

    c.location_text = snip["location"]
    c.location = eval_location(snip["location"], lk)

    if lk.title:
        t_toks = [t for t in _title_norm(lk.title).split() if t not in _TITLE_STOP]
        hay = _title_norm(c.headline + " " + snip["experience"])
        hits = sum(1 for t in t_toks if has_phrase(hay, t))
        c.title_match = bool(t_toks) and hits >= max(1, (len(t_toks) + 1) // 2)

    # score (used for ranking / ambiguity only - verdict rules below decide acceptance)
    c.score = ((40 if c.name_kind == "exact" else 34)
               + {"strong": 40, "weak": 20, "": 0}[c.company]
               + {"match_city": 15, "match_state": 8, "conflict": -30, "unknown": 0}[c.location]
               + (8 if c.title_match else 0) - 0.1 * c.position)

    located = c.location in ("match_city", "match_state")
    ok = False
    if lk.comp_full:
        if c.company == "strong":
            ok = True
        elif c.company == "weak" and (located or c.title_match):
            ok = True
        else:
            c.reasons.append("company not confirmed in search result")
    else:
        if c.location == "match_city" and c.title_match:
            ok = True
        else:
            c.reasons.append("no company given; city+title not both confirmed")

    if ok and c.location == "conflict" and not cfg.accept_location_conflict:
        ok = False
        c.reasons.append(f"location differs ({c.location_text})")

    if not ok:
        c.verdict = "review"
        return

    high = c.company == "strong" and c.name_kind == "exact" and (located or c.title_match)
    c.confidence = "High" if high else "Medium"
    if cfg.min_confidence == "high" and c.confidence != "High":
        c.verdict = "review"
        c.reasons.append("below required confidence (High)")
        return
    c.verdict = "accept"


# --------------------------------------------------------------------------- #
# SerpAPI client (caching, retries, error handling)
# --------------------------------------------------------------------------- #

class FatalSerpError(RuntimeError):
    """Invalid key / no credits - stop the whole job."""


class SerpApiClient:
    URL = "https://serpapi.com/search.json"

    def __init__(self, api_key: str, gl="us", hl="en", num=10, cache_path: Optional[str] = None,
                 timeout=60, max_retries=3, delay=(0.5, 1.5)):
        self.api_key, self.gl, self.hl, self.num = api_key, gl, hl, num
        self.timeout, self.max_retries, self.delay = timeout, max_retries, delay
        self.calls = 0                      # billable network calls this run
        self._cache: dict[str, list] = {}
        self._cache_path = cache_path
        if cache_path and Path(cache_path).exists():
            for line in Path(cache_path).read_text(encoding="utf-8").splitlines():
                try:
                    d = json.loads(line)
                    self._cache[d["k"]] = d["r"]
                except Exception:
                    continue

    def _key(self, q: str) -> str:
        return hashlib.sha1(json.dumps([q, self.gl, self.hl, self.num]).encode()).hexdigest()

    def search(self, q: str) -> list[dict]:
        k = self._key(q)
        if k in self._cache:
            return self._cache[k]
        last_err = ""
        for attempt in range(self.max_retries):
            try:
                resp = requests.get(self.URL, params={"engine": "google", "q": q, "num": self.num,
                                                      "start": 0, "gl": self.gl, "hl": self.hl,
                                                      "api_key": self.api_key}, timeout=self.timeout)
                try:
                    data = resp.json()
                except ValueError:
                    data = {}
                err = str(data.get("error") or "")
                low = err.lower()
                if "hasn't returned any results" in low or "hasn’t returned any results" in low:
                    results: list = []
                elif err:
                    if any(x in low for x in ("invalid api key", "run out of searches", "out of searches",
                                              "account", "plan")) or resp.status_code in (401, 403):
                        raise FatalSerpError(err)
                    raise RuntimeError(err)
                elif resp.status_code >= 400:
                    raise RuntimeError(f"HTTP {resp.status_code}")
                else:
                    results = data.get("organic_results") or []
                self.calls += 1
                self._cache[k] = results
                if self._cache_path:
                    with open(self._cache_path, "a", encoding="utf-8") as f:
                        f.write(json.dumps({"k": k, "r": results}) + "\n")
                time.sleep(random.uniform(*self.delay))        # "Random Delay" node
                return results
            except FatalSerpError:
                raise
            except Exception as e:                              # network / 429 / 5xx -> retry
                last_err = str(e)
                time.sleep(2 ** attempt + random.random())
        raise RuntimeError(f"SerpAPI failed after {self.max_retries} attempts: {last_err}")


# --------------------------------------------------------------------------- #
# Finder
# --------------------------------------------------------------------------- #

@dataclass
class Config:
    min_confidence: str = "medium"          # "medium" | "high"
    max_searches_per_row: int = 3
    accept_location_conflict: bool = False
    ambiguity_gap: float = 10.0             # top-2 accepted scores closer than this => ambiguous
    max_review_candidates: int = 3


def build_queries(lk: Lookup, max_n: int) -> list[str]:
    base = "site:linkedin.com/in"
    nm_full = " ".join(name_tokens(lk.full_name, allow_swap=True))
    nm_short = f"{lk.tokens[0]} {lk.tokens[-1]}"
    comp = display_company(lk.company)
    city, state = lk.city.strip(), lk.state.strip()
    q: list[str] = []
    if comp:
        if city:
            q.append(f'{base} "{nm_full}" "{comp}" "{city}"')
        q.append(f'{base} "{nm_full}" "{comp}"')
        if nm_short != nm_full:
            q.append(f'{base} "{nm_short}" "{comp}"')
        if city or state:
            q.append(f'{base} "{nm_full}" {city} {state}'.strip())
    else:
        loc = " ".join(f'"{x}"' for x in (city, state) if x)
        if lk.title:
            q.append(f'{base} "{nm_full}" "{lk.title}" {loc}'.strip())
        q.append(f'{base} "{nm_full}" {loc}'.strip())
    seen, out = set(), []
    for s in q:
        if s not in seen:
            seen.add(s)
            out.append(s)
    return out[:max_n]


class LinkedInFinder:
    def __init__(self, client, cfg: Optional[Config] = None):
        self.client, self.cfg = client, cfg or Config()

    def find(self, lk: Lookup) -> RowResult:
        if len(lk.tokens) < 2:
            return RowResult(status="Skipped - need first and last name")
        if not (lk.comp_full or lk.city_n or lk.st_full):
            return RowResult(status="Skipped - need company or city to verify")

        pool: dict[str, Candidate] = {}
        searches = 0
        accepted: list[Candidate] = []
        for q in build_queries(lk, self.cfg.max_searches_per_row):
            searches += 1
            for pos, r in enumerate(self.client.search(q)):
                c = parse_result(r, pos)
                if not c:
                    continue
                evaluate(c, lk, self.cfg)
                if c.verdict == "reject":
                    continue
                old = pool.get(c.url)
                if old is None or c.score > old.score:
                    pool[c.url] = c
            accepted = sorted((c for c in pool.values() if c.verdict == "accept"),
                              key=lambda c: -c.score)
            if accepted:
                break                       # confident answer found - don't spend more credits

        review = sorted((c for c in pool.values() if c.verdict == "review"), key=lambda c: -c.score)
        res = RowResult(searches=searches, review=review[: self.cfg.max_review_candidates])

        if not accepted:
            res.status = "Not verified - see Review sheet" if review else "Not found"
            return res
        if len(accepted) > 1 and accepted[0].score - accepted[1].score < self.cfg.ambiguity_gap:
            res.status = "Ambiguous - multiple matching profiles (left blank)"
            res.review = accepted[: self.cfg.max_review_candidates]
            return res
        best = accepted[0]
        res.url, res.confidence = best.url, best.confidence
        res.status = f"Found ({best.confidence})"
        res.matched_name, res.headline, res.evidence = best.result_name, best.headline, best.evidence()
        return res


# --------------------------------------------------------------------------- #
# Input reading / column detection
# --------------------------------------------------------------------------- #

ALIASES = {
    "full_name": ["full name", "fullname", "name", "decision maker", "decision maker name", "contact name",
                  "contact", "person", "dm name", "owner name", "owner"],
    "first_name": ["first name", "firstname", "first", "given name", "fname"],
    "last_name": ["last name", "lastname", "last", "surname", "family name", "lname"],
    "company": ["company", "company name", "organization", "organisation", "business", "business name",
                "employer", "account", "account name", "firm"],
    "city": ["city", "town"],
    "state": ["state", "state province", "province", "region", "st", "state code"],
    "country": ["country"],
    "title": ["title", "job title", "designation", "position", "role"],
    "location": ["location"],
}


def read_table(path: str) -> pd.DataFrame:
    ext = Path(path).suffix.lower()
    if ext in (".xlsx", ".xlsm", ".xls"):
        df = pd.read_excel(path, sheet_name=0, dtype=str)
    elif ext in (".csv", ".tsv", ".txt"):
        df = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                df = pd.read_csv(path, dtype=str, sep=None, engine="python", encoding=enc)
                break
            except UnicodeDecodeError:
                continue
        if df is None:
            raise ValueError("Could not decode the CSV file")
    else:
        raise ValueError(f"Unsupported file type: {ext} (use .xlsx or .csv)")
    df = df.fillna("").astype(str)
    df.columns = [str(c).strip() for c in df.columns]
    return df


def detect_columns(df: pd.DataFrame, overrides: Optional[dict] = None) -> dict:
    by_norm = {norm_text(c): c for c in df.columns}
    mapping: dict[str, str] = {}
    for key, names in ALIASES.items():
        if overrides and key in overrides:
            if overrides[key] not in df.columns:
                raise ValueError(f"Column '{overrides[key]}' not found. Available: {list(df.columns)}")
            mapping[key] = overrides[key]
            continue
        for n in names:
            if n in by_norm:
                mapping[key] = by_norm[n]
                break
    has_name = "full_name" in mapping or ("first_name" in mapping and "last_name" in mapping)
    if not has_name:
        raise ValueError(f"Could not find a name column. Columns: {list(df.columns)}. "
                         f"Use column_map={{'full_name': '<your header>'}}")
    if not any(k in mapping for k in ("company", "city", "state", "location")):
        raise ValueError("Need at least a company or city column to verify matches.")
    return mapping


def row_to_lookup(i: int, row: pd.Series, m: dict) -> Lookup:
    g = lambda k: row[m[k]].strip() if k in m else ""
    name = g("full_name") or f"{g('first_name')} {g('last_name')}".strip()
    if "full_name" in m and "first_name" in m and "last_name" in m and not g("full_name"):
        name = f"{g('first_name')} {g('last_name')}".strip()
    city, state = g("city"), g("state")
    if not city and g("location"):
        parts = [p.strip() for p in g("location").split(",")]
        city = parts[0]
        if not state and len(parts) > 1:
            state = parts[1]
    return Lookup(row_index=i, full_name=name, company=g("company"), city=city, state=state,
                  country=g("country"), title=g("title")).prepare()


# --------------------------------------------------------------------------- #
# Excel output
# --------------------------------------------------------------------------- #

def _unique(name: str, existing: list[str]) -> str:
    while name in existing:
        name += " (found)"
    return name


def write_excel(df: pd.DataFrame, results: list[RowResult], path: str, lookups: list[Lookup]) -> None:
    cols = list(df.columns)
    new = {"url": _unique("LinkedIn URL", cols), "status": _unique("Match Status", cols),
           "conf": _unique("Confidence", cols), "name": _unique("Matched Profile Name", cols),
           "head": _unique("Matched Headline", cols), "ev": _unique("Match Evidence", cols)}
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    green, amber = PatternFill("solid", fgColor="E2F0D9"), PatternFill("solid", fgColor="FFF2CC")

    headers = cols + list(new.values())
    ws.append(headers)
    for j in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=j)
        c.font, c.fill, c.alignment = bold, hdr_fill, Alignment(vertical="center", wrap_text=True)

    for i, (rec, res) in enumerate(zip(df.itertuples(index=False), results), start=2):
        ws.append(list(rec) + [res.url or None, res.status, res.confidence or None,
                               res.matched_name or None, res.headline or None, res.evidence or None])
        for j in range(1, len(headers) + 1):
            ws.cell(row=i, column=j).font = font
        url_cell = ws.cell(row=i, column=len(cols) + 1)
        if res.url:
            url_cell.hyperlink = res.url
            url_cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")
            url_cell.fill = green
        elif res.review:
            ws.cell(row=i, column=len(cols) + 2).fill = amber

    for j, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(ws.cell(row=r, column=j).value or "")) for r in range(2, min(ws.max_row, 200) + 1)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 12), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Review sheet: UNVERIFIED candidates (never written into the URL column)
    rv = wb.create_sheet("Review (unverified)")
    rv.append(["Input row", "Input name", "Input company", "Input city", "Candidate URL (UNVERIFIED)",
               "Candidate name", "Candidate headline", "Why not auto-accepted", "Evidence found"])
    for res, lk in zip(results, lookups):
        for c in res.review:
            rv.append([lk.row_index + 2, lk.full_name, lk.company, lk.city, c.url, c.result_name,
                       c.headline, "; ".join(c.reasons) or "multiple equally good matches", c.evidence()])
    for row in rv.iter_rows():
        for c in row:
            c.font = font
    for c in rv[1]:
        c.font, c.fill = bold, hdr_fill
    for j, w in enumerate([10, 24, 26, 16, 52, 24, 44, 44, 50], start=1):
        rv.column_dimensions[get_column_letter(j)].width = w
    rv.freeze_panes = "A2"

    # Summary sheet
    sm = wb.create_sheet("Summary")
    n = len(results)
    def cnt(pred): return sum(1 for r in results if pred(r))
    rows = [("Total rows", n),
            ("Found - High confidence", cnt(lambda r: r.status == "Found (High)")),
            ("Found - Medium confidence", cnt(lambda r: r.status == "Found (Medium)")),
            ("Not found (blank)", cnt(lambda r: r.status == "Not found")),
            ("Not verified - candidates on Review sheet (blank)", cnt(lambda r: r.status.startswith("Not verified"))),
            ("Ambiguous (blank)", cnt(lambda r: r.status.startswith("Ambiguous"))),
            ("Skipped / error / not processed (blank)",
             cnt(lambda r: r.status.startswith(("Skipped", "Error", "Not processed")))),
            ("SerpAPI searches used (incl. cached)", sum(r.searches for r in results))]
    sm.append(["Metric", "Value"])
    for r in rows:
        sm.append(list(r))
    for row in sm.iter_rows():
        for c in row:
            c.font = font
    for c in sm[1]:
        c.font, c.fill = bold, hdr_fill
    sm.column_dimensions["A"].width, sm.column_dimensions["B"].width = 52, 12
    sm.append([])
    sm.append(["Note: 'Found' means the search result itself showed the same name AND the company "
               "(and/or city). It cannot prove the person is *currently* employed there."])
    sm.cell(row=sm.max_row, column=1).font = Font(name="Arial", size=9, italic=True)
    wb.save(path)


# --------------------------------------------------------------------------- #
# Public API
# --------------------------------------------------------------------------- #

def process_file(input_path: str, output_path: str, api_key: Optional[str] = None,
                 column_map: Optional[dict] = None, cfg: Optional[Config] = None,
                 cache_path: Optional[str] = None, limit: Optional[int] = None,
                 client=None, progress: Optional[Callable[[int, int], None]] = None) -> dict:
    """Read CSV/XLSX, look up every row, write XLSX. Returns summary counts."""
    df = read_table(input_path)
    if limit:
        df = df.head(limit)
    mapping = detect_columns(df, column_map)
    if client is None:
        key = api_key or os.environ.get("SERPAPI_KEY")
        if not key:
            raise ValueError("Provide a SerpAPI key (api_key=... or SERPAPI_KEY env var)")
        client = SerpApiClient(key, cache_path=cache_path)
    finder = LinkedInFinder(client, cfg)

    lookups = [row_to_lookup(i, row, mapping) for i, (_, row) in enumerate(df.iterrows())]
    results: list[Optional[RowResult]] = [None] * len(lookups)
    memo: dict[str, RowResult] = {}
    fatal = ""
    for i, lk in enumerate(lookups):
        if fatal:
            results[i] = RowResult(status="Not processed (SerpAPI stopped: " + fatal + ")")
            continue
        k = lk.key()
        if k in memo and lk.tokens:                       # duplicate input row -> reuse
            results[i] = memo[k]
        else:
            try:
                results[i] = finder.find(lk)
            except FatalSerpError as e:
                fatal = str(e)
                results[i] = RowResult(status=f"Not processed (SerpAPI stopped: {fatal})")
            except Exception as e:
                results[i] = RowResult(status=f"Error: {e}")
            else:
                memo[k] = results[i]
        if progress:
            progress(i + 1, len(lookups))

    write_excel(df, results, output_path, lookups)      # always write, even after a fatal stop
    summary = {"rows": len(results),
               "found": sum(1 for r in results if r.url),
               "blank": sum(1 for r in results if not r.url),
               "searches_billed": getattr(client, "calls", None), "fatal_error": fatal or None}
    return summary


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Find LinkedIn profile URLs of decision makers via SerpAPI.")
    ap.add_argument("input")
    ap.add_argument("-o", "--output", default=None)
    ap.add_argument("--api-key", default=None, help="or set SERPAPI_KEY")
    ap.add_argument("--min-confidence", choices=["medium", "high"], default="medium")
    ap.add_argument("--max-searches", type=int, default=3, help="max SerpAPI searches per row")
    ap.add_argument("--allow-location-conflict", action="store_true")
    ap.add_argument("--cache", default="serp_cache.jsonl")
    ap.add_argument("--limit", type=int, default=None, help="only first N rows (testing)")
    ap.add_argument("--col", action="append", default=[], metavar="KEY=HEADER",
                    help="override column, e.g. --col company='Business Name' (keys: full_name, first_name, "
                         "last_name, company, city, state, country, title, location)")
    a = ap.parse_args(argv)
    overrides = dict(x.split("=", 1) for x in a.col)
    out = a.output or str(Path(a.input).with_name(Path(a.input).stem + "_linkedin.xlsx"))
    cfg = Config(min_confidence=a.min_confidence, max_searches_per_row=a.max_searches,
                 accept_location_conflict=a.allow_location_conflict)
    s = process_file(a.input, out, api_key=a.api_key, column_map=overrides or None, cfg=cfg,
                     cache_path=a.cache, limit=a.limit,
                     progress=lambda d, t: print(f"\r{d}/{t}", end="", file=sys.stderr))
    print(f"\n{json.dumps(s)}\nSaved: {out}")
    return 1 if s["fatal_error"] else 0


if __name__ == "__main__":
    sys.exit(main())
